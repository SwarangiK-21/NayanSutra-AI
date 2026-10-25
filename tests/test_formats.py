"""Document formats and OCR (offline, synthetic files): CSV/TSV, Excel, OpenDocument, RTF, PowerPoint, BMP/GIF,
scanned PDFs read with OCR, and the security checks that refuse active content in the new formats."""
from __future__ import annotations

import io
import zipfile
from datetime import datetime

import pytest
from PIL import Image, ImageDraw, ImageFont

from nyaya import filescan, ocr, parsing
from test_features import _case, _upload, _wait, client  # noqa: F401


def _write(tmp_path, name: str, data: bytes):
    p = tmp_path / name
    p.write_bytes(data)
    return p


# ── Extraction ───────────────────────────────────────────────────────────────
def test_csv_and_tsv_become_rows_with_their_column_names(tmp_path):
    csv_bytes = "Party;Amount;Due date\nAsha Rao;2,50,000;2026-03-14\nRamesh Kumar;40,000;2026-04-01\n".encode("utf-8-sig")
    p = parsing.extract(_write(tmp_path, "ledger.csv", csv_bytes))
    assert "Row 1 — Party: Asha Rao | Amount: 2,50,000 | Due date: 2026-03-14" in p.text
    assert p.meta["table"]["rows"] == 2 and p.meta["table"]["sheets"][0]["columns"] == 3
    t = parsing.extract(_write(tmp_path, "hearings.tsv", b"Date\tCourt\n2026-11-18\tSessions Court, Pune\n"))
    assert "Row 1 — Date: 2026-11-18 | Court: Sessions Court, Pune" in t.text


def test_excel_sheets_rows_dates_and_long_tables_become_pages(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payments"
    ws.append(["Instalment", "Paid on", "Amount"])
    for i in range(1, 251):
        ws.append([i, datetime(2026, 1, 1 + i % 28), 1000.0 * i])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Note"])
    ws2.append(["Cheque bounced on 14 March 2026"])
    buf = io.BytesIO()
    wb.save(buf)
    p = parsing.extract(_write(tmp_path, "payments.xlsx", buf.getvalue()))
    assert len(p.pages) == 4                                   # 250 rows → 3 pages of 100, plus the second sheet
    assert "Row 1 — Instalment: 1 | Paid on: 2026-01-02 | Amount: 1000" in p.pages[0]
    assert p.pages[2].startswith("Sheet: Payments — rows 201–250")
    assert "Cheque bounced on 14 March 2026" in p.pages[3]
    assert {s["name"] for s in p.meta["table"]["sheets"]} == {"Payments", "Notes"}


def _odf(mimetype: str, content: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", mimetype)
        z.writestr("content.xml", content)
    return buf.getvalue()


_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
       'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
       'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')


def test_opendocument_text_and_spreadsheet(tmp_path):
    odt = _odf("application/vnd.oasis.opendocument.text", f'''<office:document-content {_NS}><office:body><office:text>
        <text:h>Legal notice</text:h><text:p>Pay the outstanding rent within fifteen days.</text:p>
        <table:table><table:table-row><table:table-cell><text:p>Month</text:p></table:table-cell>
        <table:table-cell><text:p>Rent</text:p></table:table-cell></table:table-row></table:table>
        </office:text></office:body></office:document-content>''')
    t = parsing.extract(_write(tmp_path, "notice.odt", odt))
    assert t.text.splitlines() == ["Legal notice", "Pay the outstanding rent within fifteen days.", "Month | Rent"]
    ods = _odf("application/vnd.oasis.opendocument.spreadsheet", f'''<office:document-content {_NS}><office:body>
        <office:spreadsheet><table:table table:name="Dues">
        <table:table-row><table:table-cell><text:p>Tenant</text:p></table:table-cell><table:table-cell><text:p>Due</text:p></table:table-cell></table:table-row>
        <table:table-row><table:table-cell><text:p>Asha Rao</text:p></table:table-cell><table:table-cell><text:p>12000</text:p></table:table-cell>
          <table:table-cell table:number-columns-repeated="16000"/></table:table-row>
        <table:table-row table:number-rows-repeated="1048000"><table:table-cell/></table:table-row>
        </table:table></office:spreadsheet></office:body></office:document-content>''')
    s = parsing.extract(_write(tmp_path, "dues.ods", ods))     # huge repeated blank areas do not explode
    assert "Row 1 — Tenant: Asha Rao | Due: 12000" in s.text and s.meta["table"]["rows"] == 1


def test_rtf_and_powerpoint(tmp_path):
    rtf = rb"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0\pard The tenant shall vacate the premises by 30 June 2026.\par}"
    assert "vacate the premises by 30 June 2026" in parsing.extract(_write(tmp_path, "order.rtf", rtf)).text
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    s1 = deck.slides.add_slide(deck.slide_layouts[1])
    s1.shapes.title.text = "Case strategy"
    s1.placeholders[1].text = "Argue limitation under Article 54"
    s1.notes_slide.notes_text_frame.text = "Cite the sale agreement"
    s2 = deck.slides.add_slide(deck.slide_layouts[5])
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Date", "Event", "2026-03-14", "Cheque"
    buf = io.BytesIO()
    deck.save(buf)
    p = parsing.extract(_write(tmp_path, "strategy.pptx", buf.getvalue()))
    assert len(p.pages) == 2 and "Argue limitation under Article 54" in p.pages[0] and "Notes: Cite the sale agreement" in p.pages[0]
    assert "2026-03-14 | Cheque" in p.pages[1]


# ── OCR ──────────────────────────────────────────────────────────────────────
def _text_image(lines: list[str]) -> Image.Image:
    img = Image.new("RGB", (1500, 120 + 70 * len(lines)), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 34)
    except OSError:
        font = ImageFont.load_default(size=34)
    for i, ln in enumerate(lines):
        d.text((50, 50 + 70 * i), ln, font=font, fill="black")
    return img


SCAN_LINES = ["IN THE COURT OF THE SESSIONS JUDGE, PUNE", "The accused is charged under Section 318 of the BNS.",
              "Next date of hearing: 18 November 2026."]


@pytest.fixture
def rapid():
    if not ocr._rapidocr():
        pytest.skip("RapidOCR is not installed")
    ocr.engine.cache_clear()
    yield
    ocr.engine.cache_clear()


def test_scanned_pdf_is_read_with_ocr(tmp_path, rapid):
    buf = io.BytesIO()
    _text_image(SCAN_LINES).save(buf, "PDF", resolution=150)   # image-only PDF, like a scanner makes
    p = parsing.extract(_write(tmp_path, "scan.pdf", buf.getvalue()))
    text = p.text
    assert "Section 318" in text and "18 November 2026" in text and "SESSIONS JUDGE" in text
    assert p.meta["ocr"]["engine"] == "RapidOCR" and p.meta["ocr"]["pages"][0]["confidence"] > 80


def test_photos_in_every_image_format_are_read(tmp_path, rapid):
    img = _text_image(SCAN_LINES[1:2])
    for ext, fmt in ((".png", "PNG"), (".jpg", "JPEG"), (".bmp", "BMP"), (".gif", "GIF"), (".webp", "WEBP"), (".tif", "TIFF")):
        buf = io.BytesIO()
        img.save(buf, fmt)
        p = parsing.extract(_write(tmp_path, f"photo{ext}", buf.getvalue()))
        assert "Section 318" in p.text, (ext, p.text)


def test_dropped_spaces_are_restored_in_latin_text_only():
    assert ocr.respace("versusVikram charged under Section318 on 14 March2026") == \
        "versus Vikram charged under Section 318 on 14 March 2026"
    assert ocr.respace("Rs.2,50,000on 18November; hearing:18") == "Rs. 2,50,000 on 18 November; hearing: 18"
    assert ocr.respace("18th May, 10:30, McDonald, A1, s.138") == "18th May, 10:30, McDonald, A1, s.138"   # untouched
    assert ocr.respace("धारा318 के अंतर्गत") == "धारा318 के अंतर्गत"


def test_uncertain_lines_are_replaced_by_the_more_confident_reading():
    box = {"x0": 0, "x1": 100, "y0": 0, "y1": 20}
    base = [{**box, "text": "IN THE COURT", "score": 0.99}, {**box, "y0": 40, "y1": 60, "text": "R HR 318", "score": 0.6}]
    other = [{**box, "text": "garbage", "score": 0.7}, {**box, "y0": 41, "y1": 61, "text": "धारा 318", "score": 0.96},
             {**box, "y0": 90, "y1": 110, "text": "पुणे", "score": 0.95}]
    out = ocr._improve(base, other)
    assert [ln["text"] for ln in out] == ["IN THE COURT", "धारा 318", "पुणे"]   # confident lines are never replaced
    text, conf, words, low = ocr._to_text([{**box, "text": "Rent due", "score": 0.4}])
    assert text == "[?Rent] [?due]" and low == 2


# ── Upload API and security ──────────────────────────────────────────────────
def test_spreadsheet_upload_is_indexed_and_searchable(client):
    c = _case(client)
    d = _upload(client, c["id"], "rent_ledger.csv", b"Month,Rent paid\nJanuary,12000\nFebruary,0\n", "text/csv")
    assert d["status"] == "ready" and d["n_chunks"] >= 1
    hits = client.get(f"/api/search?q=February rent paid&types=document").json()["results"]
    assert any("rent_ledger" in (h.get("title") or "").lower() or "February" in (h.get("snippet") or "") for h in hits)
    f = client.get(f"/api/documents/{d['id']}/file")
    assert f.headers["content-disposition"].startswith("attachment") and f.headers["content-type"].startswith("text/csv")


def _zip(parts: dict[str, bytes | str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, v in parts.items():
            z.writestr(n, v)
    return buf.getvalue()


def _ooxml(parts: dict[str, bytes | str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        for n, v in parts.items():
            z.writestr(n, v)
    return buf.getvalue()


@pytest.mark.parametrize("name,data,why", [
    ("macro.xlsx", _ooxml({"xl/workbook.xml": "<workbook/>", "xl/vbaProject.bin": b"VBA"}), "macros"),
    ("linked.xlsx", _ooxml({"xl/workbook.xml": "<workbook/>", "xl/externalLinks/externalLink1.xml": "<x/>"}), "links"),
    ("object.pptx", _ooxml({"ppt/presentation.xml": "<p/>", "ppt/embeddings/oleObject1.bin": b"OLE"}), "embedded"),
    ("macro.ods", _zip({"content.xml": "<x/>", "Basic/Standard/Module1.xml": "<m/>"}), "macros"),
    ("object.rtf", rb"{\rtf1 Hello {\object\objemb {\objdata 0105000002000000}}}", "embedded objects"),
])
def test_active_content_in_new_formats_is_refused(tmp_path, name, data, why):
    with pytest.raises(filescan.Rejected) as exc:
        filescan.check(_write(tmp_path, name, data), "." + name.rsplit(".", 1)[1])
    assert why in str(exc.value)


def test_renamed_or_binary_files_are_refused_at_upload(client):
    c = _case(client)
    for name, data in (("ledger.xlsx", b"MZ\x90\x00 not a workbook"), ("ledger.csv", b"col\x00umn\x00binary"),
                       ("order.rtf", b"%PDF-1.4 pretending"), ("photo.gif", b"GIF00a nope"), ("run.exe", b"MZ")):
        r = client.post(f"/api/cases/{c['id']}/documents", files={"file": (name, data, "application/octet-stream")})
        assert r.status_code == 415, (name, r.status_code, r.text)
