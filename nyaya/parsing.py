"""Text extraction for PDF, Word, Excel / CSV, PowerPoint, OpenDocument, RTF, text and images (OCR), with
cleaning tuned for Indian legal documents.

PDFs go through PDFium (pypdfium2): fast, accurate word spacing, and it is the
hardened parser Chrome uses — a sensible choice for untrusted uploads.
"""
from __future__ import annotations

import re
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from . import config


class ParseError(Exception):
    """User-facing extraction problem (scanned PDF, encrypted file, ...)."""


@dataclass
class Parsed:
    pages: list[str]                 # cleaned text per page (1 entry for non-paged formats)
    meta: dict = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n".join(self.pages)

    @property
    def chars(self) -> int:
        return sum(len(p) for p in self.pages)


# ── Cleaning ─────────────────────────────────────────────────────────────────
_FOOTNOTE_REF = re.compile(r"(?<![\w\]])\d{1,2}\s*\n\s*\[")          # "1\n[(b) ..."  ->  "[(b) ..."
_FOOTNOTE_STARS = re.compile(r"\d{1,2}\s*\*\s*\*\s*\*")                # "1***" -> "***"
_FOOTNOTE_LINE = re.compile(
    r"^\s*\d{1,2}\s*\.?\s+(Subs\.|Ins\.|Omitted|Rep\.|Added|The words|Vide|Renumbered|Certain words|"
    r"Now see|See |Came into force|This Act has|Extended to|Brought into force|Cl\. |Clause \(|"
    r"Sub-section|The original|For the Statement|Amended by|Earlier|Proviso|Section \d+ (?:was|of))",
    re.I,
)
_PAGE_NUM = re.compile(r"^\s*(?:page\s*)?\d{1,4}\s*(?:of\s*\d+)?\s*$", re.I)
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_SPACES = re.compile(r"[ \t ]+")


def _norm_line(s: str) -> str:
    return re.sub(r"\d+", "#", s.strip().lower())


def clean_pages(pages: list[str]) -> list[str]:
    """Normalise unicode, strip running headers/footers, page numbers and footnote noise."""
    pages = [unicodedata.normalize("NFKC", p).replace("\r\n", "\n").replace("\r", "\n").replace("￾", "-")
             for p in pages]
    # Running headers/footers: lines (digits normalised) that recur on many pages.
    n = len(pages)
    repeated: set[str] = set()
    if n >= 4:
        counts: Counter[str] = Counter()
        for p in pages:
            lines = [ln for ln in p.split("\n") if ln.strip()]
            counts.update({_norm_line(ln) for ln in lines[:3] + lines[-3:]})
        repeated = {k for k, v in counts.items() if v >= max(3, 0.4 * n) and len(k) < 160}

    out = []
    for p in pages:
        p = _FOOTNOTE_REF.sub("[", p)
        p = _FOOTNOTE_STARS.sub("***", p)
        kept = []
        for ln in p.split("\n"):
            s = _SPACES.sub(" ", ln).strip()
            if not s or _PAGE_NUM.match(s) or s == "IndiaCode" or _norm_line(s) in repeated:
                if not s:
                    kept.append("")
                continue
            if _FOOTNOTE_LINE.match(s):
                continue
            kept.append(s)
        text = "\n".join(kept)
        text = _HYPHEN_BREAK.sub(r"\1\2", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        out.append(text)
    return out


# ── Extractors ───────────────────────────────────────────────────────────────
def _pdf(path: Path, progress=None) -> Parsed:
    import pypdfium2 as pdfium

    try:
        doc = pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as exc:
        msg = str(exc).lower()
        if "password" in msg:
            raise ParseError("This PDF is password-protected. Remove the password and upload again.") from exc
        raise ParseError("This file could not be read as a PDF (it may be damaged).") from exc
    meta: dict = {}
    try:
        if len(doc) > config.MAX_PDF_PAGES:
            raise ParseError(f"PDF has {len(doc)} pages; the limit is {config.MAX_PDF_PAGES}.")
        raw = []
        for i in range(len(doc)):
            page = doc[i]
            tp = page.get_textpage()
            raw.append(tp.get_text_range())
            tp.close()
            page.close()
        # Pages without a text layer (scans) are OCR'd when an engine is installed.
        scanned = [i for i, t in enumerate(raw) if len(t.strip()) < 25]
        if scanned:
            meta.update(_ocr_pdf_pages(doc, scanned, raw, all_scanned=len(scanned) == len(raw), progress=progress))
    finally:
        doc.close()
    from .quality import printed_numbers
    meta["printed_page_numbers"] = printed_numbers(raw)   # read before cleaning removes page-number lines
    pages = clean_pages(raw)
    total = sum(len(p) for p in pages)
    if total < max(200, 25 * len(pages)) and not meta.get("ocr"):
        raise ParseError(
            "No readable text found — this looks like a scanned PDF and no OCR engine is available "
            "(install the app's requirements: RapidOCR, or Tesseract), so the text cannot be read reliably."
        )
    return Parsed(pages=pages, meta={"pages": len(pages), **meta})


def _ocr_pdf_pages(doc, idx: list[int], raw: list[str], all_scanned: bool, progress=None) -> dict:
    from . import ocr

    if not ocr.engine():
        return {"ocr_needed_pages": [i + 1 for i in idx]} if not all_scanned else {}
    if len(idx) > config.OCR_MAX_PAGES:
        raise ParseError(f"{len(idx)} scanned pages need OCR; the limit is {config.OCR_MAX_PAGES} "
                         "(OCR_MAX_PAGES in .env).")
    per_page = []
    for k, i in enumerate(idx):
        if progress:   # OCR is slow: report each page (this is also where a cancel takes effect)
            progress(k, len(idx), f"Reading scanned page {i + 1} of {len(raw)} (OCR)")
        page = doc[i]
        try:
            img = page.render(scale=300 / 72).to_pil()  # ~300 dpi, what Tesseract is tuned for
        finally:
            page.close()
        try:
            res = ocr.ocr_image(img)
        except RuntimeError as exc:
            raise ParseError(f"OCR failed on page {i + 1}: {exc}.") from exc
        raw[i] = res["text"]
        per_page.append({"page": i + 1, "confidence": res["confidence"], "low_words": res["low_words"],
                         "words": res["words"]})
    words = sum(p["words"] for p in per_page) or 1
    return {"ocr": {"engine": ocr.engine()["engine"], "langs": ocr.engine()["langs"], "pages": per_page,
                    "low_conf_ratio": round(sum(p["low_words"] for p in per_page) / words, 3)}}


def _image(path: Path, progress=None) -> Parsed:
    from . import ocr

    if not ocr.engine():
        raise ParseError("Image documents need OCR, and no OCR engine is available "
                         "(install the app's requirements: RapidOCR, or Tesseract).")
    try:
        img = ocr.open_image(path)
    except Exception as exc:  # PIL raises many types, including DecompressionBombError
        raise ParseError("This image could not be opened (damaged, unsupported or too large).") from exc
    frames = []
    try:
        n = getattr(img, "n_frames", 1)  # multi-page TIFF
        for f in range(min(n, config.OCR_MAX_PAGES)):
            img.seek(f)
            frames.append(img.copy())
    except EOFError:
        pass
    pages, per_page = [], []
    for i, frame in enumerate(frames or [img], 1):
        if progress:
            progress(i - 1, len(frames or [img]), f"Reading image page {i} (OCR)")
        try:
            res = ocr.ocr_image(frame)
        except RuntimeError as exc:
            raise ParseError(f"OCR failed on page {i}: {exc}.") from exc
        pages.append(res["text"])
        per_page.append({"page": i, "confidence": res["confidence"], "low_words": res["low_words"], "words": res["words"]})
    words = sum(p["words"] for p in per_page) or 1
    return Parsed(pages=clean_pages(pages), meta={"pages": len(pages), "ocr": {
        "engine": ocr.engine()["engine"], "langs": ocr.engine()["langs"], "pages": per_page,
        "low_conf_ratio": round(sum(p["low_words"] for p in per_page) / words, 3)}})


def _doc(path: Path) -> Parsed:
    """Legacy Word 97-2003 (.doc) via antiword when available."""
    import shutil
    import subprocess

    exe = shutil.which("antiword")
    if not exe:
        raise ParseError("Legacy .doc files need the 'antiword' tool, which is not installed. "
                         "Open the file in Word and save it as .docx, then upload again.")
    try:
        out = subprocess.run([exe, "-w", "0", str(path)], capture_output=True, timeout=60)
    except subprocess.TimeoutExpired as exc:
        raise ParseError("Reading the .doc file timed out.") from exc
    if out.returncode != 0:
        raise ParseError("This .doc file could not be read. Save it as .docx and upload again.")
    return Parsed(pages=clean_pages([out.stdout.decode("utf-8", errors="replace")]), meta={"pages": 1})


def _docx(path: Path) -> Parsed:
    # Zip-bomb guard before handing the archive to python-docx.
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > 300 * 1024 * 1024:
            raise ParseError("DOCX expands to an unreasonable size; refusing to open it.")
        if "word/document.xml" not in z.namelist():
            raise ParseError("Not a valid Word document.")
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return Parsed(pages=clean_pages(["\n".join(parts)]), meta={"pages": 1})


def _text(path: Path) -> Parsed:
    return Parsed(pages=clean_pages([_decode(path.read_bytes())]), meta={"pages": 1})


# ── Spreadsheets and other office formats ────────────────────────────────────
MAX_TABLE_ROWS = 50_000          # per file
MAX_TABLE_CELLS = 2_000_000
ROWS_PER_PAGE = 100              # a "page" of a sheet = 100 rows, so citations can point at a row range


def _cell(v) -> str:
    """Cell value as text: dates as ISO, whole-number floats without '.0'."""
    from datetime import date, datetime, time

    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="minutes").removesuffix(" 00:00")
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else f"{v:g}"
    return re.sub(r"\s+", " ", str(v)).strip()


def _table_pages(sheets: list[tuple[str, list[list]]]) -> Parsed:
    """Rows become lines 'Row n — Header: value | Header: value' (the first non-empty row is the header), so a
    question about one entry finds that row with its column names."""
    pages, cells, rows_total, out_sheets = [], 0, 0, []
    for name, rows in sheets:
        rows = [[_cell(v) for v in r] for r in rows]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        header = [h or f"Column {i + 1}" for i, h in enumerate(rows[0] + [""] * (width - len(rows[0])))]
        body = rows[1:] if len(rows) > 1 else rows
        out_sheets.append({"name": name, "rows": len(body), "columns": width})
        for start in range(0, len(body), ROWS_PER_PAGE):
            lines = [f"Sheet: {name} — rows {start + 1}–{min(start + ROWS_PER_PAGE, len(body))}",
                     "Columns: " + " | ".join(header)]
            for i, r in enumerate(body[start:start + ROWS_PER_PAGE], start + 1):
                vals = [f"{header[j]}: {v}" for j, v in enumerate(r) if v]
                lines.append(f"Row {i} — " + " | ".join(vals))
            pages.append("\n".join(lines))
        rows_total += len(body)
        cells += sum(len(r) for r in rows)
        if rows_total > MAX_TABLE_ROWS or cells > MAX_TABLE_CELLS:
            raise ParseError(f"The spreadsheet is too large (over {MAX_TABLE_ROWS:,} rows or {MAX_TABLE_CELLS:,} cells). "
                             "Split it, or export the part you need.")
    if not pages:
        raise ParseError("The spreadsheet has no data.")
    return Parsed(pages=pages, meta={"pages": len(pages), "table": {"sheets": out_sheets, "rows": rows_total}})


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _csv(path: Path) -> Parsed:
    import csv

    text = _decode(path.read_bytes())
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    csv.field_size_limit(1_000_000)
    rows = []
    for i, row in enumerate(csv.reader(text.splitlines(), dialect)):
        if i > MAX_TABLE_ROWS:
            raise ParseError(f"The file has more than {MAX_TABLE_ROWS:,} rows. Split it and upload the parts.")
        rows.append(row)
    return _table_pages([(path.stem, rows)])


def _zip_guard(path: Path, required: str, label: str) -> list[str]:
    try:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 300 * 1024 * 1024:
                raise ParseError(f"The {label} file expands to an unreasonable size; refusing to open it.")
            names = z.namelist()
    except zipfile.BadZipFile:
        raise ParseError(f"The {label} file is damaged.") from None
    if required not in names:
        raise ParseError(f"Not a valid {label} file.")
    return names


def _xlsx(path: Path) -> Parsed:
    _zip_guard(path, "xl/workbook.xml", "Excel")
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)   # cached values; formulas never run
    except Exception as exc:
        raise ParseError("This Excel file could not be read. Save it again as .xlsx and upload.") from exc
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_TABLE_ROWS:
                    raise ParseError(f"A sheet has more than {MAX_TABLE_ROWS:,} rows. Split it and upload the parts.")
                rows.append(list(r))
            sheets.append((ws.title, rows))
    finally:
        wb.close()
    return _table_pages(sheets)


def _xls(path: Path) -> Parsed:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise ParseError("This old-format Excel file could not be read. Save it as .xlsx and upload again.") from exc
    try:
        sheets = []
        for sh in book.sheets():
            if sh.nrows > MAX_TABLE_ROWS:
                raise ParseError(f"A sheet has more than {MAX_TABLE_ROWS:,} rows. Split it and upload the parts.")
            rows = []
            for r in range(sh.nrows):
                vals = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            vals.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                            continue
                        except (ValueError, OverflowError):
                            pass
                    vals.append(cell.value)
                rows.append(vals)
            sheets.append((sh.name, rows))
    finally:
        book.release_resources()
    return _table_pages(sheets)


_ODF = {"office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
        "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
        "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}


def _odf_root(path: Path, label: str):
    from defusedxml import ElementTree

    _zip_guard(path, "content.xml", label)
    with zipfile.ZipFile(path) as z:
        try:
            return ElementTree.fromstring(z.read("content.xml"))
        except Exception as exc:
            raise ParseError(f"This {label} file could not be read.") from exc


def _odf_text(el) -> str:
    return re.sub(r"\s+", " ", "".join(el.itertext())).strip()


def _ods(path: Path) -> Parsed:
    root = _odf_root(path, "OpenDocument spreadsheet")
    t = "{%s}" % _ODF["table"]
    sheets = []
    for tab in root.iter(t + "table"):
        rows = []
        for row in tab.iter(t + "table-row"):
            vals = []
            for cell in row:
                if cell.tag not in (t + "table-cell", t + "covered-table-cell"):
                    continue
                n = min(int(cell.get(t + "number-columns-repeated", "1") or 1), 1000)   # trailing blanks repeat a lot
                vals += [_odf_text(cell)] * n
            while vals and not vals[-1]:
                vals.pop()
            rep = min(int(row.get(t + "number-rows-repeated", "1") or 1), 1 if not any(vals) else 1000)
            rows += [vals] * rep
            if len(rows) > MAX_TABLE_ROWS:
                raise ParseError(f"A sheet has more than {MAX_TABLE_ROWS:,} rows. Split it and upload the parts.")
        sheets.append((tab.get(t + "name", "Sheet"), rows))
    return _table_pages(sheets)


def _odt(path: Path) -> Parsed:
    root = _odf_root(path, "OpenDocument text")
    tx, tb = "{%s}" % _ODF["text"], "{%s}" % _ODF["table"]
    parts: list[str] = []

    def walk(el) -> None:   # paragraphs and headings in document order; a table becomes one line per row
        for child in el:
            if child.tag == tb + "table":
                for row in child.iter(tb + "table-row"):
                    cells = [_odf_text(c) for c in row if c.tag == tb + "table-cell"]
                    if any(cells):
                        parts.append(" | ".join(c for c in cells if c))
            elif child.tag in (tx + "h", tx + "p"):
                s = _odf_text(child)
                if s:
                    parts.append(s)
            else:
                walk(child)
    walk(root)
    return Parsed(pages=clean_pages(["\n".join(parts)]), meta={"pages": 1})


def _rtf(path: Path) -> Parsed:
    from striprtf.striprtf import rtf_to_text

    raw = path.read_bytes()
    try:
        text = rtf_to_text(raw.decode("latin-1"), errors="ignore")
    except Exception as exc:
        raise ParseError("This RTF file could not be read. Save it as .docx or PDF and upload again.") from exc
    return Parsed(pages=clean_pages([text]), meta={"pages": 1})


def _pptx(path: Path) -> Parsed:
    _zip_guard(path, "ppt/presentation.xml", "PowerPoint")
    from pptx import Presentation

    try:
        deck = Presentation(str(path))
    except Exception as exc:
        raise ParseError("This PowerPoint file could not be read. Save it as PDF and upload again.") from exc
    pages = []
    for slide in deck.slides:
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        pages.append("\n".join(parts))
    return Parsed(pages=clean_pages(pages), meta={"pages": len(pages)})


EXTRACTORS = {".pdf": _pdf, ".docx": _docx, ".doc": _doc, ".txt": _text, ".md": _text, ".rtf": _rtf, ".odt": _odt,
              ".csv": _csv, ".tsv": _csv, ".xlsx": _xlsx, ".xls": _xls, ".ods": _ods, ".pptx": _pptx,
              ".png": _image, ".jpg": _image, ".jpeg": _image, ".tif": _image, ".tiff": _image, ".webp": _image,
              ".bmp": _image, ".gif": _image}


def extract(path: Path, progress=None) -> Parsed:
    """`progress(done, total, text)` is called for each OCR page (scanned PDFs, images)."""
    from . import langdetect

    fn = EXTRACTORS.get(path.suffix.lower())
    if fn is None:
        raise ParseError(f"Unsupported file type {path.suffix}. Use PDF, Word, Excel, CSV, PowerPoint, OpenDocument, "
                         "RTF, text or an image.")
    parsed = fn(path, progress) if fn in (_pdf, _image) else fn(path)
    if parsed.chars < 20:
        raise ParseError("The document appears to be empty.")
    parsed.meta["language"] = langdetect.detect(parsed.text)
    return parsed
